"""固定结构的质量异常调查 Excel 导出。"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from api.app.domain import InvestigationCase
from api.app.domain.ai import EvidenceGroundedInvestigationResult

_TITLE_FILL = PatternFill("solid", fgColor="1F4F79")
_SECTION_FILL = PatternFill("solid", fgColor="DCE9F3")
_HEADER_FILL = PatternFill("solid", fgColor="2E708D")
_SUBTLE_FILL = PatternFill("solid", fgColor="F4F7FA")
_WHITE_FONT = Font(color="FFFFFF", bold=True)
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_LABEL_FONT = Font(color="243B53", bold=True)
_THIN_SIDE = Side(style="thin", color="CDD8E3")
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
_WRAP_TOP = Alignment(vertical="top", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_MAX_CELL_TEXT = 30_000

_STATE_LABELS = {
    "created": "已创建",
    "mapping_required": "待确认字段",
    "ready": "证据已就绪",
    "investigating": "正在分析",
    "waiting_for_user": "等待人工判断",
    "completed": "已完成",
    "failed": "失败",
}
_HYPOTHESIS_STATUS_LABELS = {
    "candidate": "候选",
    "under_verification": "验证中",
    "confirmed": "已确认",
    "rejected": "已排除",
}
_ACTION_TYPE_LABELS = {
    "containment": "围堵",
    "verification": "验证",
    "corrective": "纠正",
}


def _safe_text(value: object) -> str:
    """清洗单元格文本，并阻止来自上传数据或模型文本的公式注入。"""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        text = value.isoformat()
    elif isinstance(value, (dict, list, tuple, set)):
        text = json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False, default=str)
    else:
        text = str(value)

    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(text) > _MAX_CELL_TEXT:
        text = text[: _MAX_CELL_TEXT - 1] + "…"
    stripped = text.lstrip()
    if stripped and stripped[0] in "=+-@":
        text = "'" + text
    return text


def _cell_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    return _safe_text(value)


def _source_names(source_refs: Iterable[str]) -> str:
    return "、".join(Path(value).name for value in source_refs)


def _prepare_sheet(sheet: Worksheet, *, landscape: bool = True) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A4"
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def _write_title(sheet: Worksheet, title: str, *, column_count: int) -> None:
    last_column = get_column_letter(max(column_count, 1))
    sheet.merge_cells(f"A1:{last_column}1")
    cell = sheet["A1"]
    cell.value = _safe_text(title)
    cell.fill = _TITLE_FILL
    cell.font = Font(color="FFFFFF", bold=True, size=16)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28


def _apply_table_style(
    sheet: Worksheet,
    *,
    header_row: int,
    first_row: int,
    last_row: int,
    column_count: int,
) -> None:
    for cell in sheet[header_row]:
        if cell.column > column_count:
            break
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER

    for row in sheet.iter_rows(
        min_row=first_row,
        max_row=max(last_row, first_row),
        min_col=1,
        max_col=column_count,
    ):
        for cell in row:
            cell.alignment = _WRAP_TOP
            cell.border = _THIN_BORDER

    if last_row >= header_row:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(column_count)}{last_row}"


def _auto_size_columns(sheet: Worksheet, *, max_width: int = 56) -> None:
    for column_index in range(1, sheet.max_column + 1):
        width = 10
        for cell in sheet.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=1,
            max_row=sheet.max_row,
        ):
            for item in cell:
                value = "" if item.value is None else str(item.value)
                longest_line = max((len(line) for line in value.splitlines()), default=0)
                width = max(width, min(longest_line + 2, max_width))
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _write_table(
    sheet: Worksheet,
    *,
    title: str,
    headers: list[str],
    rows: Iterable[Iterable[object]],
) -> None:
    _prepare_sheet(sheet)
    _write_title(sheet, title, column_count=len(headers))
    header_row = 3
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column_index, value=header)

    current_row = header_row + 1
    for row_values in rows:
        for column_index, value in enumerate(row_values, start=1):
            sheet.cell(row=current_row, column=column_index, value=_cell_value(value))
        current_row += 1

    if current_row == header_row + 1:
        sheet.cell(row=current_row, column=1, value="无记录")
        current_row += 1

    _apply_table_style(
        sheet,
        header_row=header_row,
        first_row=header_row + 1,
        last_row=current_row - 1,
        column_count=len(headers),
    )
    _auto_size_columns(sheet)


def _write_summary(
    sheet: Worksheet,
    *,
    case: InvestigationCase,
    ai_result: EvidenceGroundedInvestigationResult | None,
) -> None:
    _prepare_sheet(sheet, landscape=False)
    _write_title(sheet, "质量异常调查报告", column_count=4)

    metadata = [
        ("案件编号", case.case_id),
        ("调查问题", case.question),
        ("案件状态", _STATE_LABELS.get(case.state, case.state)),
        ("质量特征", case.scope.quality_feature or "未识别"),
        ("调查开始", case.scope.start_at),
        ("调查结束", case.scope.end_at),
        ("源文件", _source_names(case.source_refs)),
        ("证据数量", len(case.evidence)),
        ("候选根因数量", len(case.hypotheses)),
        ("动作数量", len(case.actions)),
        ("创建时间", case.created_at),
        ("更新时间", case.updated_at),
    ]

    row = 3
    for label, value in metadata:
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=_cell_value(value))
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        label_cell = sheet.cell(row=row, column=1)
        label_cell.fill = _SECTION_FILL
        label_cell.font = _LABEL_FONT
        label_cell.alignment = _WRAP_TOP
        label_cell.border = _THIN_BORDER
        for column_index in range(2, 5):
            value_cell = sheet.cell(row=row, column=column_index)
            value_cell.alignment = _WRAP_TOP
            value_cell.border = _THIN_BORDER
        row += 1

    sections: list[tuple[str, list[str]]] = [
        ("数据缺口", case.missing_data),
        ("最终结论", [case.conclusion] if case.conclusion else []),
    ]
    if ai_result is not None:
        sections.insert(0, ("AI 证据发现", [item.statement for item in ai_result.findings]))

    for section_title, values in sections:
        row += 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        section_cell = sheet.cell(row=row, column=1, value=section_title)
        section_cell.fill = _HEADER_FILL
        section_cell.font = _WHITE_FONT
        section_cell.alignment = Alignment(vertical="center")
        row += 1
        if not values:
            values = ["无"]
        for index, value in enumerate(values, start=1):
            sheet.cell(row=row, column=1, value=index)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            sheet.cell(row=row, column=2, value=_safe_text(value))
            for column_index in range(1, 5):
                cell = sheet.cell(row=row, column=column_index)
                cell.alignment = _WRAP_TOP
                cell.border = _THIN_BORDER
                if index % 2 == 0:
                    cell.fill = _SUBTLE_FILL
            row += 1

    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 28


def _evidence_rows(case: InvestigationCase) -> Iterable[list[object]]:
    for item in case.evidence:
        yield [
            item.id,
            item.evidence_type,
            item.title,
            item.statement,
            item.confidence,
            item.sample_size,
            item.origin,
            item.filters,
            item.metrics,
            _source_names(item.source_refs),
            item.created_at,
        ]


def _hypothesis_rows(case: InvestigationCase) -> Iterable[list[object]]:
    for item in case.hypotheses:
        decision = item.decision
        yield [
            item.id,
            _HYPOTHESIS_STATUS_LABELS.get(item.status, item.status),
            item.confidence,
            item.statement,
            "、".join(item.supporting_evidence_ids),
            "、".join(item.contradicting_evidence_ids),
            "\n".join(item.verification_actions),
            decision.actor_id if decision else "",
            decision.note if decision else "",
            decision.decided_at if decision else "",
        ]


def _action_rows(case: InvestigationCase) -> Iterable[list[object]]:
    for item in case.actions:
        yield [
            item.id,
            _ACTION_TYPE_LABELS.get(item.action_type, item.action_type),
            item.status,
            item.title,
            item.rationale,
            "、".join(item.related_hypothesis_ids),
            "、".join(item.supporting_evidence_ids),
            item.owner,
            item.due_at,
            item.completed_at,
        ]


def _ai_plan_rows(
    ai_result: EvidenceGroundedInvestigationResult | None,
) -> Iterable[list[object]]:
    if ai_result is None:
        return []
    return [
        [index, step.tool_name, step.reason, step.arguments]
        for index, step in enumerate(ai_result.plan.steps, start=1)
    ]


def render_investigation_report(
    *,
    case: InvestigationCase,
    output_path: Path,
    ai_result: EvidenceGroundedInvestigationResult | None = None,
) -> Path:
    """使用固定表结构导出案件；模型不能选择模板、单元格或样式。"""
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("investigation export path must end with .xlsx")

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "调查摘要"
    _write_summary(summary_sheet, case=case, ai_result=ai_result)

    evidence_sheet = workbook.create_sheet("证据")
    _write_table(
        evidence_sheet,
        title="调查证据",
        headers=[
            "证据编号",
            "证据类型",
            "标题",
            "证据陈述",
            "可信度",
            "样本量",
            "来源类型",
            "过滤条件",
            "指标",
            "源文件",
            "生成时间",
        ],
        rows=_evidence_rows(case),
    )

    hypothesis_sheet = workbook.create_sheet("候选根因")
    _write_table(
        hypothesis_sheet,
        title="候选根因与人工决定",
        headers=[
            "假设编号",
            "状态",
            "可信度",
            "候选原因",
            "支持证据",
            "反向证据",
            "验证动作",
            "人工操作者",
            "人工说明",
            "决定时间",
        ],
        rows=_hypothesis_rows(case),
    )

    action_sheet = workbook.create_sheet("动作")
    _write_table(
        action_sheet,
        title="围堵、验证与纠正动作",
        headers=[
            "动作编号",
            "动作类型",
            "状态",
            "动作",
            "理由",
            "相关假设",
            "支持证据",
            "负责人",
            "到期时间",
            "完成时间",
        ],
        rows=_action_rows(case),
    )

    plan_sheet = workbook.create_sheet("AI调查计划")
    _write_table(
        plan_sheet,
        title="AI 白名单工具计划",
        headers=["序号", "工具", "选择理由", "参数"],
        rows=_ai_plan_rows(ai_result),
    )
    if ai_result is not None:
        plan_sheet["A2"] = _safe_text(ai_result.plan.interpreted_question)
        plan_sheet.merge_cells("A2:D2")
        plan_sheet["A2"].alignment = _WRAP_TOP

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
