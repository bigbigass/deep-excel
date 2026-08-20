from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from api.app.domain import (
    Hypothesis,
    HypothesisDecision,
    InvestigationAction,
    InvestigationCase,
    InvestigationFinding,
    InvestigationPlan,
)
from api.app.domain.ai import EvidenceGroundedInvestigationResult
from api.app.main import app
from api.app.services.investigation import cases
from api.app.services.investigation.baseline import run_baseline_investigation
from api.app.services.investigation.export import render_investigation_report
from api.app.services.investigation.repository import InvestigationRepository


DEMO_DATA_PATH = Path("sample_data/investigation_demo.csv")
client = TestClient(app)


def _build_case(*, question: str = "为什么外径不良率升高？") -> InvestigationCase:
    frame = pd.read_csv(DEMO_DATA_PATH, parse_dates=["measured_at"])
    base = run_baseline_investigation(
        frame,
        question=question,
        case_id="CASE-EXPORT-1",
        source_refs=[str(DEMO_DATA_PATH)],
    )
    evidence_id = base.evidence[0].id
    decision = HypothesisDecision(
        outcome="confirmed",
        actor_id="quality-engineer-01",
        note="完成现场检查和复测后确认。",
    )
    hypothesis = Hypothesis(
        id="H-EXPORT-1",
        statement="M02 的 4 号型腔定位偏移可能是候选因素。",
        supporting_evidence_ids=[evidence_id],
        confidence="medium",
        status="confirmed",
        verification_actions=["检查型腔定位状态"],
        decision=decision,
    )
    action = InvestigationAction(
        id="A-EXPORT-1",
        action_type="verification",
        title="检查型腔定位状态",
        rationale="验证候选假设。",
        status="completed",
        related_hypothesis_ids=[hypothesis.id],
        supporting_evidence_ids=[evidence_id],
        completed_at=decision.decided_at,
    )
    payload = base.model_dump()
    payload.update(
        {
            "question": question,
            "state": "waiting_for_user",
            "hypotheses": [hypothesis],
            "actions": [action],
        }
    )
    return InvestigationCase.model_validate(payload)


def _build_ai_result(case: InvestigationCase) -> EvidenceGroundedInvestigationResult:
    evidence_id = case.evidence[0].id
    return EvidenceGroundedInvestigationResult(
        plan=InvestigationPlan(
            interpreted_question="定位外径不良升高的生产关联因素。",
            steps=[],
            missing_information=[],
        ),
        findings=[
            InvestigationFinding(
                id="F-AI-001",
                finding_type="fact",
                statement="数据质量检查已经完成。",
                evidence_ids=[evidence_id],
                confidence="high",
            )
        ],
        missing_data=[],
    )


def test_render_investigation_report_contains_evidence_and_human_decision(tmp_path: Path) -> None:
    case = _build_case()
    output_path = tmp_path / "investigation.xlsx"

    rendered = render_investigation_report(
        case=case,
        ai_result=_build_ai_result(case),
        output_path=output_path,
    )

    assert rendered == output_path
    workbook = openpyxl.load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["调查摘要", "证据", "候选根因", "动作", "AI调查计划"]
    assert workbook["调查摘要"]["B3"].value == "CASE-EXPORT-1"
    assert workbook["调查摘要"]["B4"].value == "为什么外径不良率升高？"
    assert workbook["证据"]["A4"].value == case.evidence[0].id
    assert workbook["候选根因"]["A4"].value == "H-EXPORT-1"
    assert workbook["候选根因"]["B4"].value == "已确认"
    assert workbook["候选根因"]["H4"].value == "quality-engineer-01"
    assert workbook["候选根因"]["I4"].value == "完成现场检查和复测后确认。"
    assert workbook["动作"]["A4"].value == "A-EXPORT-1"
    assert workbook["AI调查计划"]["A2"].value == "定位外径不良升高的生产关联因素。"


def test_investigation_export_neutralizes_excel_formula_injection(tmp_path: Path) -> None:
    case = _build_case(question='=HYPERLINK("https://example.test","open")')
    output_path = tmp_path / "safe.xlsx"

    render_investigation_report(case=case, output_path=output_path)

    workbook = openpyxl.load_workbook(output_path, data_only=False)
    question_cell = workbook["调查摘要"]["B4"]
    assert question_cell.data_type != "f"
    assert question_cell.value == '\'=HYPERLINK("https://example.test","open")'


def test_investigation_export_requires_xlsx_path(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="must end with .xlsx"):
        render_investigation_report(
            case=_build_case(),
            output_path=tmp_path / "investigation.csv",
        )


def _create_ready_case(repository: InvestigationRepository) -> str:
    payload = cases.create_investigation(
        question="为什么外径不良率升高？",
        file_name=DEMO_DATA_PATH.name,
        content=DEMO_DATA_PATH.read_bytes(),
        repository=repository,
        start_background=False,
    )
    case_id = payload["case_id"]
    initial = repository.load(case_id)
    source_path = repository.resolve_source_ref(case_id, initial.source_refs[0])
    completed = cases.run_investigation(case_id, source_path, repository=repository)
    assert completed.state == "ready", completed.error
    return case_id


def test_investigation_export_api_returns_xlsx(monkeypatch, tmp_path: Path) -> None:
    repository = InvestigationRepository(tmp_path / "investigations")
    monkeypatch.setattr(cases, "_DEFAULT_REPOSITORY", repository)
    case_id = _create_ready_case(repository)

    response = client.get(f"/api/v1/investigations/{case_id}/export")

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"{case_id}-investigation.xlsx" in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(BytesIO(response.content))
    assert workbook["调查摘要"]["B3"].value == case_id
    assert workbook["证据"].max_row > 4


def test_investigation_export_api_returns_not_found(monkeypatch, tmp_path: Path) -> None:
    repository = InvestigationRepository(tmp_path / "investigations")
    monkeypatch.setattr(cases, "_DEFAULT_REPOSITORY", repository)

    response = client.get("/api/v1/investigations/CASE-NOT-FOUND/export")

    assert response.status_code == 404
